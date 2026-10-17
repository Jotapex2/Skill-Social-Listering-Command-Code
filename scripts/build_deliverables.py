"""Genera nubes de palabras, volumen, gráficos por hora, temas e informe (MD + DOCX).

Uso:
    python scripts/build_deliverables.py

Requiere haber corrido antes `build_md_db.py` (db/menciones.csv).
"""
from __future__ import annotations

import json
import os
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from wordcloud import WordCloud

import pos

ROOT = Path(__file__).resolve().parent.parent
CONFIG = ROOT / "config" / "keywords.json"
DB_CSV = ROOT / "db" / "menciones.csv"
OUT = ROOT / "output"

STOP = set("""a al algo algunas algunos ante antes aqui asi aun aunque como con contra cual
cuales cuando de del desde donde dos e el ella ellas ellos en entre era eran es esa esas ese
eso esos esta estaba estan estar estas este esto estos fin fue fueron ha hace hacia han hasta
hay la las le les lo los mas me mi mientras mis mucho muchos muy nada ni no nos nosotros o os
otra otras otro otros para pero poco por porque primer que quien quienes se sea segun ser si
sin sobre sois solo son su sus tambien tanto te tener tiene tienen todo todos tu tus un una
unas uno unos usted vosotros y ya yo rt via https http www com instagram reel video videos
q dia dias hoy ayer manana ahora vamos ver gente lugar parte cosa cosas hacer dijo
""".split())

# Se rellena en main() con las keywords del tema (no son concepto de la conversación).
NOISE: set[str] = set()

PALETTE = [(31, 56, 100), (44, 104, 119), (56, 113, 184), (100, 76, 153), (73, 80, 91)]
ACCENT = (226, 100, 46)
BLUE = (56, 113, 184)
GRAY_L = (211, 216, 223)
INK = (18, 26, 42)
GRAY = (95, 103, 115)


def _resolve_fonts() -> tuple[str | None, str | None]:
    """Busca fuentes en rutas típicas (Windows/Linux/macOS). None -> fuente por defecto de PIL."""
    bold = ["C:/Windows/Fonts/arialbd.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
            "/Library/Fonts/Arial Bold.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"]
    regular = ["C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
               "/Library/Fonts/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf"]
    fb = next((p for p in bold if os.path.exists(p)), None)
    fr = next((p for p in regular if os.path.exists(p)), None)
    return fb, fr


FONT_B, FONT_R = _resolve_fonts()


def _get_font(bold: bool, size: int):
    path = FONT_B if bold else FONT_R
    if path:
        return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def norm(s: str) -> str:
    s = unicodedata.normalize("NFD", str(s or "").casefold())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def parse_num(v) -> int:
    if v is None:
        return 0
    s = str(v).strip().replace(",", "")
    if not s:
        return 0
    m = re.match(r"([\d.]+)\s*([KMB])?(?![A-Za-z])", s, re.I)
    if not m:
        return 0
    try:
        n = float(m.group(1))
    except ValueError:
        return 0
    return int(n * {"k": 1e3, "m": 1e6, "b": 1e9}.get((m.group(2) or "").lower(), 1))


def tokenize(text: str, drop_subject: bool = False) -> list[str]:
    t = re.sub(r"https?://\S+|www\.\S+", " ", str(text or ""))
    t = re.sub(r"[@#][\w.]+", " ", t)
    words = re.findall(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]{3,}", t)
    out = []
    for w in words:
        n = norm(w)
        if n in STOP or len(n) < 3:
            continue
        if drop_subject and n in NOISE:
            continue
        out.append(n)
    return out


def ngrams(tokens: list[str], n: int) -> Counter:
    c = Counter()
    for i in range(len(tokens) - n + 1):
        g = tokens[i : i + n]
        if len(set(g)) == 1:
            continue
        c[" ".join(g)] += 1
    return c


def make_cloud(path: Path, freq: dict[str, int], polarity: str | None = None) -> int:
    if not freq:
        Image.new("RGB", (2000, 900), "white").save(path)
        return 0
    wc = WordCloud(
        width=2000, height=900, background_color="white", collocations=False,
        prefer_horizontal=0.985, relative_scaling=0.40, min_font_size=13,
        max_font_size=240, max_words=90, margin=1, random_state=42,
        font_path=FONT_B, normalize_plurals=False,
    ).generate_from_frequencies(freq)
    if polarity in ("positivo", "negativo"):
        lo, hi = min(freq.values()), max(freq.values())

        def color(word, **kw):
            t = (freq[word] - lo) / (hi - lo) if hi > lo else 1
            if polarity == "positivo":
                return (int(27 + 129 * (1 - t)), int(111 + 101 * (1 - t)), int(53 + 111 * (1 - t)))
            return (int(145 + 82 * (1 - t)), int(22 + 128 * (1 - t)), int(34 + 120 * (1 - t)))

        wc.recolor(color_func=color, random_state=42)
    else:
        colors = {w: PALETTE[i % len(PALETTE)] for i, (w, _) in enumerate(wc.words_.items())}
        wc.recolor(color_func=lambda word, **kw: colors.get(word, PALETTE[0]), random_state=42)
    wc.to_image().convert("RGB").save(path)
    return len(wc.words_)


def make_bar_chart(path: Path, horas: list[tuple[str, int]], titulo: str) -> bool:
    """Gráfico de barras de posteos por hora. True si se generó con datos."""
    img = Image.new("RGB", (2000, 900), "white")
    d = ImageDraw.Draw(img)
    d.text((90, 60), titulo, font=_get_font(True, 56), fill=INK, anchor="la")
    if not horas:
        d.text((90, 170), "Sin datos de hora para esta plataforma.", font=_get_font(False, 40),
               fill=GRAY, anchor="la")
        img.save(path)
        return False
    horas = sorted(horas, key=lambda x: x[0])
    mx = max(v for _, v in horas) or 1
    x0, x1 = 150, 1940
    y0, y1 = 250, 780
    slot = (x1 - x0) / len(horas)
    bw = min(int(slot * 0.6), 120)
    d.line([(x0 - 20, y1), (x1 + 20, y1)], fill=GRAY_L, width=3)
    for i, (hora, v) in enumerate(horas):
        cx = int(x0 + slot * (i + 0.5))
        h = int((y1 - y0) * v / mx)
        top = y1 - h
        peak = v == mx
        d.rounded_rectangle([cx - bw // 2, top, cx + bw // 2, y1], radius=8,
                            fill=ACCENT if peak else BLUE)
        d.text((cx, top - 14), str(v), font=_get_font(True, 30 if peak else 26),
               fill=ACCENT if peak else INK, anchor="ms")
        d.text((cx, y1 + 18), str(hora), font=_get_font(False, 28), fill=GRAY, anchor="mt")
    d.text((1940, 800), "posteos por hora", font=_get_font(False, 30), fill=GRAY, anchor="ra")
    img.save(path)
    return True


def fmt(n) -> str:
    return f"{n:,}".replace(",", ".")


def at(v) -> str:
    v = str(v or "").strip()
    return v if v.startswith("@") else "@" + v


def _horas(df: pd.DataFrame) -> list[tuple[str, int]]:
    if df.empty or "hora" not in df:
        return []
    sub = df[df["hora"] != ""]
    if sub.empty:
        return []
    g = sub.groupby("hora").size().reset_index(name="posts")
    return [(str(r["hora"]), int(r["posts"])) for _, r in g.iterrows()]


def main() -> None:
    global NOISE
    cfg = json.loads(CONFIG.read_text(encoding="utf-8"))
    tema = cfg.get("tema", "conversación")
    fecha = cfg["fecha_objetivo"]
    kws = cfg["keywords"]
    NOISE = {norm(k) for k in kws}
    slug = re.sub(r"[^0-9A-Za-z]+", "_", tema).strip("_").lower() or "informe"
    OUT.mkdir(parents=True, exist_ok=True)

    narrativa_path = ROOT / "config" / "narrativa_conversacion.md"
    narrativa = narrativa_path.read_text(encoding="utf-8").strip() if narrativa_path.exists() else ""
    if "Escribe aquí la narrativa" in narrativa:  # no inyectar la plantilla sin editar
        narrativa = ""
    if narrativa:
        narrativa = re.sub(r"^#+\s+.*\n?", "", narrativa, count=1).strip()

    df = pd.read_csv(DB_CSV, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    if df.empty:
        raise SystemExit("db/menciones.csv está vacío: corre primero la recolección.")
    for col in ("likes", "comentarios", "reposts", "vistas"):
        if col not in df:
            df[col] = ""
        df[col + "_n"] = df[col].map(parse_num)
    df["engagement"] = df["likes_n"] + df["comentarios_n"] + df["reposts_n"]
    df["hora"] = df["hora_local"].str.slice(11, 13)

    comments = pd.DataFrame()
    cpath = ROOT / "db" / "comentarios.csv"
    if cpath.exists():
        comments = pd.read_csv(cpath, dtype=str, keep_default_na=False, encoding="utf-8-sig")
        if "likes" in comments:
            comments["likes_n"] = comments["likes"].map(parse_num)

    total = len(df)
    por_plataforma = df["plataforma"].value_counts().to_dict()

    kw_counts = Counter()
    for _, r in df.iterrows():
        t = norm(str(r.get("texto", "")))
        for k in kws:
            if norm(k) in t:
                kw_counts[k] += 1

    vol_rows = []
    for p, n in por_plataforma.items():
        sub = df[df["plataforma"] == p]
        vol_rows.append({
            "plataforma": p, "posts": n,
            "likes": int(sub["likes_n"].sum()), "comentarios": int(sub["comentarios_n"].sum()),
            "reposts": int(sub["reposts_n"].sum()), "vistas": int(sub["vistas_n"].sum()),
            "engagement": int(sub["engagement"].sum()),
            "autores_unicos": sub["autor"].nunique(),
        })
    pd.DataFrame(vol_rows).to_csv(OUT / "volumen_plataforma.csv", index=False, encoding="utf-8")

    horas_gen = _horas(df)
    pd.DataFrame(horas_gen, columns=["hora", "posts"]).to_csv(
        OUT / "volumen_por_hora.csv", index=False, encoding="utf-8")

    pd.DataFrame(sorted(kw_counts.items(), key=lambda x: -x[1]),
                 columns=["keyword", "posts"]).to_csv(
        OUT / "volumen_keywords.csv", index=False, encoding="utf-8")

    top_autores = (df.groupby(["plataforma", "autor"])
                   .agg(posts=("id", "count"), engagement=("engagement", "sum"))
                   .reset_index().sort_values("posts", ascending=False).head(25))
    top_autores.to_csv(OUT / "top_autores.csv", index=False, encoding="utf-8")

    influencers = []
    for p in por_plataforma:
        sub = df[df["plataforma"] == p]
        agg = (sub.groupby("autor")
               .agg(posts=("id", "count"), likes=("likes_n", "sum"),
                    comentarios=("comentarios_n", "sum"), engagement=("engagement", "sum"))
               .reset_index()
               .query("autor != ''")
               .sort_values(["posts", "engagement"], ascending=False)
               .head(5))
        for _, r in agg.iterrows():
            influencers.append({
                "plataforma": p, "autor": r["autor"], "posts": int(r["posts"]),
                "likes": int(r["likes"]), "comentarios": int(r["comentarios"]),
                "engagement": int(r["engagement"]),
            })
    infl_df = pd.DataFrame(influencers)
    if not infl_df.empty:
        infl_df.to_csv(OUT / "top_influencers.csv", index=False, encoding="utf-8")

    top_posts = df.sort_values("engagement", ascending=False).head(20)
    top_posts[["plataforma", "autor", "hora_local", "likes", "comentarios", "reposts",
               "url", "texto"]].to_csv(OUT / "top_posts.csv", index=False, encoding="utf-8")

    # --- Nubes generales (sustantivos permitidos) ---
    ctextos = list(comments["comentario"]) if not comments.empty else []
    gen_tokens = [t for txt in df["texto"] for t in tokenize(txt)]
    gen_tokens += [t for txt in ctextos for t in tokenize(txt)]
    freq_gen = {w: c for w, c in Counter(gen_tokens).most_common(120)}
    cloud_words = make_cloud(OUT / "nube_palabras.png", freq_gen)
    pd.DataFrame([{"palabra": w, "frecuencia": c} for w, c in Counter(gen_tokens).most_common(150)]
                 ).to_csv(OUT / "palabras_general.csv", index=False, encoding="utf-8")
    for p in por_plataforma:
        toks = [t for txt in df[df["plataforma"] == p]["texto"] for t in tokenize(txt)]
        if not comments.empty and "plataforma" in comments:
            toks += [t for txt in comments[comments["plataforma"] == p]["comentario"] for t in tokenize(txt)]
        make_cloud(OUT / f"nube_palabras_{p}.png", {w: c for w, c in Counter(toks).most_common(90)})

    # --- Nubes positivas/negativas: SOLO verbos y adjetivos, sin nombres propios ---
    allowed = pos.content_words(list(df["texto"]) + ctextos)
    print(f"POS: {len(allowed)} verbos/adjetivos permitidos; spaCy={'sí' if pos.is_available() else 'no (heurístico)'}")

    def filtrado(text: str) -> list[str]:
        return [t for t in tokenize(text, drop_subject=True) if t in allowed]

    sent_posts: dict[str, int] = {}
    pos_tokens, neg_tokens = [], []
    for _, r in df.iterrows():
        lab = r.get("tono", "") or "sin_clasificar"
        sent_posts[lab] = sent_posts.get(lab, 0) + 1
        toks = filtrado(r.get("texto", ""))
        if lab == "positivo":
            pos_tokens += toks
        elif lab == "negativo":
            neg_tokens += toks
    sent_comments: dict[str, int] = {}
    if not comments.empty and "sentimiento" in comments:
        for lab, group in comments.groupby("sentimiento"):
            sent_comments[lab] = len(group)
            toks = [t for txt in group["comentario"] for t in filtrado(txt)]
            if lab == "positivo":
                pos_tokens += toks
            elif lab == "negativo":
                neg_tokens += toks
    freq_pos = {w: c for w, c in Counter(pos_tokens).most_common(90)}
    freq_neg = {w: c for w, c in Counter(neg_tokens).most_common(90)}
    pos_words = make_cloud(OUT / "nube_palabras_positivas.png", freq_pos, "positivo")
    neg_words = make_cloud(OUT / "nube_palabras_negativas.png", freq_neg, "negativo")
    pd.DataFrame([{"palabra": w, "frecuencia": c} for w, c in Counter(pos_tokens).most_common(150)]
                 ).to_csv(OUT / "palabras_positivas.csv", index=False, encoding="utf-8")
    pd.DataFrame([{"palabra": w, "frecuencia": c} for w, c in Counter(neg_tokens).most_common(150)]
                 ).to_csv(OUT / "palabras_negativas.csv", index=False, encoding="utf-8")
    pd.DataFrame(
        [{"alcance": "posts", "sentimiento": k, "n": v} for k, v in sorted(sent_posts.items())]
        + [{"alcance": "comentarios", "sentimiento": k, "n": v} for k, v in sorted(sent_comments.items())]
    ).to_csv(OUT / "sentimiento.csv", index=False, encoding="utf-8")

    top_comentarios = pd.DataFrame()
    if not comments.empty:
        top_comentarios = comments.sort_values("likes_n", ascending=False).head(25)
        top_comentarios.to_csv(OUT / "top_comentarios.csv", index=False, encoding="utf-8")
        (comments.groupby(["plataforma", "post_url"]).size()
         .reset_index(name="comentarios_recolectados")
         .sort_values("comentarios_recolectados", ascending=False)
         .to_csv(OUT / "comentarios_por_post.csv", index=False, encoding="utf-8"))

    # --- Gráficos de posteos por hora (general + por red) ---
    graficos = [("Posts por hora — general", "grafico_por_hora.png",
                 make_bar_chart(OUT / "grafico_por_hora.png", horas_gen,
                                f"Posteos por hora — {tema} ({fecha})"))]
    horas_por_plataforma: dict[str, list[tuple[str, int]]] = {}
    for p in por_plataforma:
        hs = _horas(df[df["plataforma"] == p])
        horas_por_plataforma[p] = hs
        nombre = "X" if p == "x" else "Instagram"
        graficos.append((f"Posts por hora — {nombre}", f"grafico_por_hora_{p}.png",
                         make_bar_chart(OUT / f"grafico_por_hora_{p}.png", hs,
                                        f"Posteos por hora — {nombre} ({fecha})")))
    ig_sin_hora = "instagram" in por_plataforma and not horas_por_plataforma.get("instagram")

    # --- Temas y frases frecuentes ---
    temas = defaultdict(lambda: {"posts": 0, "plataformas": Counter(), "ejemplos": [],
                                 "likes": 0, "comentarios": 0})
    for _, r in df.iterrows():
        t = r.get("tema", "") or "otros"
        d = temas[t]
        d["posts"] += 1
        d["plataformas"][r["plataforma"]] += 1
        d["likes"] += int(r["likes_n"])
        d["comentarios"] += int(r["comentarios_n"])
        if len(d["ejemplos"]) < 4 and r.get("url"):
            d["ejemplos"].append((r.get("autor", ""), str(r.get("texto", ""))[:180], r["url"]))

    bigrams = ngrams(gen_tokens, 2)
    trigrams = ngrams(gen_tokens, 3)
    frases = [(g, c) for g, c in (bigrams + trigrams).most_common(25) if c >= 3]

    tema_rows = [{
        "tema": t, "posts": d["posts"],
        "plataformas": ";".join(f"{k}:{v}" for k, v in d["plataformas"].most_common()),
        "likes": d["likes"], "comentarios": d["comentarios"],
    } for t, d in sorted(temas.items(), key=lambda x: -x[1]["posts"])]
    pd.DataFrame(tema_rows).to_csv(OUT / "temas.csv", index=False, encoding="utf-8")
    pd.DataFrame(frases, columns=["frase", "menciones"]).to_csv(
        OUT / "frases_frecuentes.csv", index=False, encoding="utf-8")

    # --- Informe MD ---
    plataforma_txt = " y ".join(f"**{p}** ({fmt(n)} posts)" for p, n in por_plataforma.items())
    md = []
    md.append(f"# Social listening — {tema} ({fecha})")
    md.append("")
    md.append("## Resumen ejecutivo")
    md.append("")
    md.append(
        f"Se recolectaron **{fmt(total)} publicaciones** del {fecha} en {plataforma_txt}, "
        f"a partir de las keywords {', '.join(kws)}. "
        f"La conversación suma {fmt(int(df['likes_n'].sum()))} likes, "
        f"{fmt(int(df['comentarios_n'].sum()))} comentarios y "
        f"{fmt(int(df['reposts_n'].sum()))} reposteos. "
        + (f"Se recolectó el texto de **{fmt(len(ctextos))} comentarios**."
           if not comments.empty else "No se recolectaron comentarios.")
    )
    md.append("")
    if narrativa:
        md.append("## Resumen de la conversación")
        md.append("")
        md.append(narrativa)
        md.append("")
    md.append("## Metodología")
    md.append("")
    md.append(
        "Recolección **solo por scraping de navegador** (`agent-browser` sobre un perfil aislado): "
        "X por búsqueda (`x.com/search`, operador `since/until`) e Instagram por *keyword* "
        "(no hashtag). Unidad de conteo: una publicación única, deduplicada por URL canónica. "
        "Fechas convertidas a America/Santiago antes de asignar el día."
    )
    md.append("")
    md.append("## Volumen de conversación")
    md.append("")
    md.append("| Plataforma | Posts | Likes | Comentarios | Reposts | Engagement | Autores únicos |")
    md.append("|---|---|---|---|---|---|---|")
    for r in vol_rows:
        md.append(f"| {r['plataforma']} | {r['posts']} | {fmt(r['likes'])} | "
                  f"{fmt(r['comentarios'])} | {fmt(r['reposts'])} | {fmt(r['engagement'])} | "
                  f"{r['autores_unicos']} |")
    md.append("")
    md.append("**Posts por keyword**")
    md.append("")
    md.append("| Keyword | Posts |")
    md.append("|---|---|")
    for k, v in kw_counts.most_common():
        md.append(f"| {k} | {v} |")
    md.append("")

    md.append("## Posteos por hora")
    md.append("")
    if horas_gen:
        peak = max(horas_gen, key=lambda x: x[1])
        ptot = sum(v for _, v in horas_gen)
        md.append(f"El peak de publicación se concentra a las **{peak[0]}:00 h** "
                  f"con {peak[1]} publicaciones.")
        md.append("")
        md.append("| Hora | Posts | % del total |")
        md.append("|---|---|---|")
        for hora, v in horas_gen:
            pct = f"{100 * v / ptot:.1f}".replace(".", ",")
            md.append(f"| {hora} h | {v} | {pct}% |")
        md.append(f"| **Total** | **{fmt(ptot)}** | **100%** |")
        md.append("")
    md.append(f"![Posts por hora — general](grafico_por_hora.png)")
    md.append("")
    for nombre, archivo, ok in graficos[1:]:
        if ok:
            md.append(f"![{nombre}]({archivo})")
            md.append("")

    md.append("## Top 5 influencers por plataforma")
    md.append("")
    md.append("Ranking por **cantidad de posteos** sobre el tema en la fecha; desempata el engagement.")
    md.append("")
    for p in por_plataforma:
        nombre = "X" if p == "x" else "Instagram"
        md.append(f"### {nombre}")
        md.append("")
        md.append("| # | Cuenta | Posts | Likes | Comentarios | Engagement |")
        md.append("|---|---|---|---|---|---|")
        sub = infl_df[infl_df["plataforma"] == p] if not infl_df.empty else pd.DataFrame()
        if sub.empty:
            md.append("| — | (sin datos) | | | | |")
        for i, (_, r) in enumerate(sub.iterrows(), 1):
            md.append(f"| {i} | {r['autor']} | {r['posts']} | {fmt(r['likes'])} | "
                      f"{fmt(r['comentarios'])} | {fmt(r['engagement'])} |")
        md.append("")
    md.append("## Nube de palabras")
    md.append("")
    md.append(f"![Nube de palabras](nube_palabras.png) — {cloud_words} términos (top mostrado).")
    md.append("")
    md.append("## Conceptos positivos y negativos")
    md.append("")
    md.append(
        "Clasificación por **léxico español conservador** (sin señal → `sin_clasificar`). "
        "Las nubes de conceptos incluyen **solo verbos y adjetivos** (sin nombres propios ni "
        "sustantivos institucionales), calculadas sobre los textos clasificados. "
        "Posts: " + ", ".join(f"{k}: {fmt(v)}" for k, v in sorted(sent_posts.items()))
        + ("." if not sent_comments else ". Comentarios: "
           + ", ".join(f"{k}: {fmt(v)}" for k, v in sorted(sent_comments.items())) + ".")
    )
    md.append("")
    md.append(f"![Conceptos positivos](nube_palabras_positivas.png) — {pos_words} términos")
    md.append("")
    md.append(f"![Conceptos negativos](nube_palabras_negativas.png) — {neg_words} términos")
    md.append("")
    md.append("## Conversaciones más frecuentes")
    md.append("")
    for row in tema_rows:
        d = temas[row["tema"]]
        md.append(f"### {row['tema']} — {row['posts']} posts")
        md.append("")
        for autor, texto, url in d["ejemplos"]:
            md.append(f"- {texto} — {at(autor)} ([enlace]({url}))")
        md.append("")
    if frases:
        md.append("**Frases más repetidas**")
        md.append("")
        for g, c in frases[:15]:
            md.append(f"- \"{g}\" ({c})")
        md.append("")
    md.append("## Comentarios")
    md.append("")
    if comments.empty:
        md.append("No se recolectaron comentarios en esta corrida.")
    else:
        porc = comments["plataforma"].value_counts().to_dict()
        md.append(
            f"Se recolectaron **{fmt(len(comments))} comentarios** "
            f"({', '.join(f'{k}: {fmt(v)}' for k, v in porc.items())}), "
            "extraídos de los posts recolectados. La nube de palabras incluye posts + comentarios."
        )
        md.append("")
        md.append("| Plataforma | Autor | Likes | Comentario | Enlace |")
        md.append("|---|---|---|---|---|")
        for _, c in top_comentarios.iterrows():
            txt = re.sub(r"\|", "¦", str(c["comentario"]))[:160].replace("\n", " ")
            md.append(f"| {c['plataforma']} | {c['autor']} | {c['likes']} | {txt} | "
                      f"[post]({c['post_url']}) |")
    md.append("")
    md.append("## Publicaciones con más engagement")
    md.append("")
    md.append("| Plataforma | Autor | Hora | Likes | Com. | Reposts | Texto | Enlace |")
    md.append("|---|---|---|---|---|---|---|---|")
    for _, r in top_posts.iterrows():
        txt = re.sub(r"\|", "¦", str(r["texto"]))[:110].replace("\n", " ")
        md.append(f"| {r['plataforma']} | {at(r['autor'])} | {r['hora_local']} | {r['likes']} | "
                  f"{r['comentarios']} | {r['reposts']} | {txt} | [link]({r['url']}) |")
    md.append("")
    md.append("## Sentimiento: positivo, negativo y neutro")
    md.append("")
    md.append(
        "Clasificación por léxico español conservador. **Neutro** son los textos sin señal léxica "
        "(`sin_clasificar`): se cuentan como neutros y se declara, no se reparten entre positivo y "
        "negativo."
    )
    md.append("")
    md.append("| Alcance | Positivo | Negativo | Neutro | Total |")
    md.append("|---|---|---|---|---|")
    for alcance, dic in (("Posteos", sent_posts), ("Comentarios", sent_comments)):
        n_pos = dic.get("positivo", 0)
        n_neg = dic.get("negativo", 0)
        n_neu = dic.get("sin_clasificar", 0) + dic.get("neutral", 0)
        md.append(f"| {alcance} | {fmt(n_pos)} | {fmt(n_neg)} | {fmt(n_neu)} | {fmt(n_pos + n_neg + n_neu)} |")
    md.append("")
    md.append("## Límites")
    md.append("")
    limites = [
        "- Cobertura parcial: los buscadores de IG/X no son deterministas ni exhaustivos.",
        "- Recolección solo por scraping de navegador: X puede pedir login o aplicar rate limit.",
        "- En Instagram sólo se accede a contenido público visible con sesión iniciada.",
        "- Las métricas se toman de lo que muestra la interfaz y pueden variar con el tiempo.",
        "- El sentimiento es orientativo: clasificador por léxico español, conservador; los textos "
        "sin señal quedan como `sin_clasificar`.",
        "- Las nubes de conceptos positivos/negativos muestran solo verbos y adjetivos (top N); el "
        "detalle está en `palabras_positivas.csv` / `palabras_negativas.csv`.",
    ]
    if ig_sin_hora:
        limites.append("- Instagram no entregó hora fiable, así que su gráfico por hora se omitió: "
                       "se reporta el general y el de X.")
    md.append("\n".join(limites))
    md.append("")
    (OUT / f"informe_{slug}_{fecha}.md").write_text("\n".join(md) + "\n", encoding="utf-8")

    write_docx(OUT / f"informe_{slug}_{fecha}.docx", vol_rows, tema_rows, top_posts, tema, fecha,
               total, kws, frases, top_comentarios, infl_df, narrativa, sent_posts, sent_comments,
               horas_gen, graficos)
    print(json.dumps({
        "total": total, "por_plataforma": por_plataforma, "temas": [t["tema"] for t in tema_rows],
        "nube_terminos": cloud_words, "pos_spacy": pos.is_available(),
        "graficos": [g[1] for g in graficos if g[2]], "output": str(OUT),
    }, ensure_ascii=False, indent=2))


def write_docx(path, vol_rows, tema_rows, top_posts, tema, fecha, total, kws, frases,
               top_comentarios=None, infl_df=None, narrativa="", sent_posts=None,
               sent_comments=None, horas_gen=None, graficos=None):
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_heading(f"Social listening — {tema} ({fecha})", level=0)
    doc.add_paragraph(f"{total} publicaciones recolectadas del {fecha}. Keywords: {', '.join(kws)}.")
    if narrativa:
        doc.add_heading("Resumen de la conversación", level=1)
        for block in [b.strip() for b in narrativa.split("\n\n") if b.strip()]:
            txt = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", block)
            txt = re.sub(r"\s+", " ", txt.replace("**", "").replace("\n", " ")).strip()
            doc.add_paragraph(txt)
    doc.add_heading("Volumen de conversación", level=1)
    t = doc.add_table(rows=1, cols=7)
    t.style = "Table Grid"
    for i, h in enumerate(["Plataforma", "Posts", "Likes", "Comentarios", "Reposts",
                           "Engagement", "Autores únicos"]):
        t.rows[0].cells[i].text = h
    for r in vol_rows:
        c = t.add_row().cells
        for i, k in enumerate(["plataforma", "posts", "likes", "comentarios", "reposts",
                               "engagement", "autores_unicos"]):
            c[i].text = fmt(r[k]) if k != "plataforma" else r[k]

    if horas_gen:
        doc.add_heading("Posteos por hora", level=1)
        ptot = sum(v for _, v in horas_gen)
        ht = doc.add_table(rows=1, cols=3)
        ht.style = "Table Grid"
        for i, h in enumerate(["Hora", "Posts", "% del total"]):
            ht.rows[0].cells[i].text = h
        for hora, v in horas_gen:
            c = ht.add_row().cells
            c[0].text = f"{hora} h"
            c[1].text = str(v)
            c[2].text = f"{100 * v / ptot:.1f}%".replace(".", ",")
    if graficos:
        for nombre, archivo, ok in graficos:
            if ok:
                doc.add_heading(nombre, level=2)
                doc.add_picture(str(OUT / archivo), width=Inches(6.3))

    if infl_df is not None and not infl_df.empty:
        doc.add_heading("Top 5 influencers por plataforma (por cantidad de posteos)", level=1)
        for p in infl_df["plataforma"].unique():
            doc.add_heading("X" if p == "x" else "Instagram", level=2)
            it = doc.add_table(rows=1, cols=5)
            it.style = "Table Grid"
            for i, h in enumerate(["#", "Cuenta", "Posts", "Likes", "Engagement"]):
                it.rows[0].cells[i].text = h
            for i, (_, r) in enumerate(infl_df[infl_df["plataforma"] == p].iterrows(), 1):
                c = it.add_row().cells
                c[0].text = str(i)
                c[1].text = str(r["autor"])
                c[2].text = str(r["posts"])
                c[3].text = str(r["likes"])
                c[4].text = str(r["engagement"])

    for nombre, etiqueta in (("nube_palabras.png", "Nube de palabras"),
                             ("nube_palabras_positivas.png", "Conceptos positivos"),
                             ("nube_palabras_negativas.png", "Conceptos negativos")):
        p = OUT / nombre
        if p.exists():
            doc.add_heading(etiqueta, level=1)
            doc.add_picture(str(p), width=Inches(6.3))
    doc.add_heading("Conversaciones más frecuentes", level=1)
    for row in tema_rows:
        doc.add_heading(f"{row['tema']} — {row['posts']} posts", level=2)
        doc.add_paragraph(f"Plataformas: {row['plataformas']}")
    if frases:
        doc.add_paragraph("Frases más repetidas: " + "; ".join(f'"{g}" ({c})' for g, c in frases[:10]))
    if top_comentarios is not None and not top_comentarios.empty:
        doc.add_heading("Comentarios", level=1)
        ct = doc.add_table(rows=1, cols=4)
        ct.style = "Table Grid"
        for i, h in enumerate(["Plataforma", "Autor", "Likes", "Comentario"]):
            ct.rows[0].cells[i].text = h
        for _, c in top_comentarios.head(15).iterrows():
            row = ct.add_row().cells
            row[0].text = str(c["plataforma"])
            row[1].text = str(c["autor"])
            row[2].text = str(c["likes"])
            row[3].text = re.sub(r"\s+", " ", str(c["comentario"]))[:200]
    doc.add_heading("Publicaciones con más engagement", level=1)
    tt = doc.add_table(rows=1, cols=4)
    tt.style = "Table Grid"
    for i, h in enumerate(["Plataforma", "Autor", "Likes", "Enlace"]):
        tt.rows[0].cells[i].text = h
    for _, r in top_posts.head(12).iterrows():
        c = tt.add_row().cells
        c[0].text, c[1].text, c[2].text = r["plataforma"], at(r["autor"]), str(r["likes"])
        c[3].text = str(r["url"])
    if sent_posts is not None:
        doc.add_heading("Sentimiento: positivo, negativo y neutro", level=1)
        doc.add_paragraph(
            "Clasificación por léxico español conservador. Neutro = textos sin señal léxica "
            "(`sin_clasificar`). Las nubes de conceptos incluyen solo verbos y adjetivos."
        )
        st = doc.add_table(rows=1, cols=5)
        st.style = "Table Grid"
        for i, h in enumerate(["Alcance", "Positivo", "Negativo", "Neutro", "Total"]):
            st.rows[0].cells[i].text = h
        for alcance, dic in (("Posteos", sent_posts), ("Comentarios", sent_comments or {})):
            pos = dic.get("positivo", 0)
            neg = dic.get("negativo", 0)
            neu = dic.get("sin_clasificar", 0) + dic.get("neutral", 0)
            row = st.add_row().cells
            row[0].text = alcance
            row[1].text = fmt(pos)
            row[2].text = fmt(neg)
            row[3].text = fmt(neu)
            row[4].text = fmt(pos + neg + neu)
    try:
        doc.save(str(path))
    except PermissionError:
        print(f"[aviso] no se pudo guardar {path} (¿está abierto en Word?). El MD sí quedó actualizado.")


if __name__ == "__main__":
    main()
